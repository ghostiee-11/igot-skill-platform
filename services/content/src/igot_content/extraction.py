import io,ipaddress,re,socket
from pathlib import Path
from urllib.parse import quote_plus,urlparse
import asyncio
import httpx
from bs4 import BeautifulSoup
class ExtractionError(ValueError):pass
def extract_document(filename:str,data:bytes)->str:
    extension=Path(filename or "").suffix.lower()
    if extension==".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if extension==".pptx":
        from pptx import Presentation
        presentation=Presentation(io.BytesIO(data));parts=[]
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:parts.append(shape.text_frame.text)
            if slide.has_notes_slide:parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)
    if extension==".docx":
        from docx import Document
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
    if extension in {".txt",".md",".vtt",".srt"}:
        value=data.decode("utf-8",errors="replace")
        return clean_transcript(value) if extension in {".vtt",".srt"} else value
    raise ExtractionError("Upload PDF, PPTX, DOCX, TXT, MD, VTT, or SRT learning material")
def normalize_text(value:str)->str:return re.sub(r"\n{3,}","\n\n",re.sub(r"[ \t]+"," ",value)).strip()
async def extract_web_text(uri:str)->tuple[str,dict]:
    parsed=urlparse(uri)
    if parsed.scheme not in {"http","https"}:raise ExtractionError("only http and https sources are supported")
    if not parsed.hostname:raise ExtractionError("source URL must include a hostname")
    try:
        addresses=await asyncio.to_thread(socket.getaddrinfo,parsed.hostname,parsed.port or (443 if parsed.scheme=="https" else 80),type=socket.SOCK_STREAM)
    except socket.gaierror as exc:raise ExtractionError("source hostname could not be resolved") from exc
    for address in {item[4][0] for item in addresses}:
        if not ipaddress.ip_address(address).is_global:raise ExtractionError("private or reserved source addresses are not allowed")
    async with httpx.AsyncClient(timeout=20,follow_redirects=True,headers={"User-Agent":"iGOT-content-ingestor/1.0"}) as client:
        res=await client.get(uri);res.raise_for_status()
    content_type=res.headers.get("content-type","")
    if "html" not in content_type and not content_type.startswith("text/"):raise ExtractionError(f"unsupported content type: {content_type}")
    soup=BeautifulSoup(res.text,"html.parser")
    for node in soup(["script","style","noscript","nav","footer"]):node.decompose()
    title=soup.title.get_text(" ",strip=True) if soup.title else "";text=normalize_text(soup.get_text("\n"));return text,{"title":title,"content_type":content_type,"bytes":len(res.content)}
def discovery_links(query:str,limit:int=10)->list[dict]:
    q=quote_plus(query);providers=[("YouTube",f"https://www.youtube.com/results?search_query={q}"),("DIKSHA",f"https://diksha.gov.in/explore-course?searchText={q}"),("iGOT Karmayogi",f"https://igotkarmayogi.gov.in/#/explore?search={q}")];return [{"provider":p,"title":query,"url":u,"kind":"catalog_search"} for p,u in providers][:limit]

def clean_transcript(raw_text:str)->str:
    value=re.sub(r"WEBVTT.*?\n","",raw_text,flags=re.I)
    value=re.sub(r"\d{1,2}:\d{2}(?::\d{2})?(?:[.,]\d{3})?\s*-->\s*\d{1,2}:\d{2}(?::\d{2})?(?:[.,]\d{3})?","",value)
    value=re.sub(r"^\s*\d+\s*$","",value,flags=re.M)
    value=re.sub(r"\[.*?\]|\(.*?\)","",value)
    return normalize_text(value)

def chunk_text(value:str,size:int=1500,overlap:int=200)->list[dict]:
    if not value:return []
    chunks=[];start=0
    while start<len(value):
        end=min(start+size,len(value))
        if end<len(value):
            boundary=max(value.rfind("\n\n",start,end),value.rfind(". ",start,end))
            if boundary>start:end=boundary+1
        text=value[start:end].strip()
        if text:chunks.append({"chunk_index":len(chunks),"text":text,"token_count":max(1,len(text)//4),"start_char":start,"end_char":end,"metadata":{"chunk_index":len(chunks)}})
        if end>=len(value):break
        start=max(end-overlap,start+1)
    return chunks
